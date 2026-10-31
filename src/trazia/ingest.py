import hashlib
import io
import zipfile
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook
from PIL import Image
from pydantic import Field

from .model import Document, Model, Source

MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 100
IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".tif", ".tiff")
IMAGE_DPI = 200  # el mismo del OCR: el PDF de una imagen se reconoce a su tamaño original
# Firma de los primeros bytes por extensión: la extensión sola no dice qué es el archivo.
SIGNATURES = {
    ".pdf": (b"%PDF-",),
    ".xlsx": (b"PK\x03\x04",),
    ".png": (b"\x89PNG\r\n\x1a\n",),
    ".jpg": (b"\xff\xd8\xff",),
    ".jpeg": (b"\xff\xd8\xff",),
    ".tif": (b"II*\x00", b"MM\x00*"),
    ".tiff": (b"II*\x00", b"MM\x00*"),
}


def _check_signature(data: bytes, suffix: str) -> None:
    signatures = SIGNATURES.get(suffix)
    if signatures is None:
        return
    # Un PDF puede llevar basura antes de la cabecera; los lectores la aceptan en el primer KB.
    head = data[:1024] if suffix == ".pdf" else data[:16]
    found = any(s in head for s in signatures) if suffix == ".pdf" else head.startswith(signatures)
    if not found:
        raise ValueError(
            f"No es un {suffix.lstrip('.').upper()} válido: no corresponde a la extensión"
        )


def image_to_pdf(data: bytes) -> bytes:
    """Una imagen (PNG, JPEG o TIFF de varias páginas) como PDF escaneado, para el OCR y la visión.

    Cada píxel queda a `IMAGE_DPI`: al renderizar el PDF para el OCR sale la imagen original.
    """
    with Image.open(io.BytesIO(data)) as image:
        frames = []
        for index in range(getattr(image, "n_frames", 1)):
            if index >= MAX_PAGES:
                raise ValueError("Límite de la base: 100 páginas por documento")
            image.seek(index)
            frames.append(image.convert("RGB"))
    buffer = io.BytesIO()
    frames[0].save(buffer, "PDF", resolution=IMAGE_DPI, save_all=True, append_images=frames[1:])
    return buffer.getvalue()


class EvidenceLine(Model):
    text: str
    source: Source


class Bundle(Model):
    document: Document
    media_type: str
    lines: list[EvidenceLine] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)


def ingest_bytes(data: bytes, filename: str, document_id: str, revision: str) -> Bundle:
    if not data or len(data) > MAX_BYTES:
        raise ValueError("El documento debe ocupar entre 1 byte y 25 MB")
    suffix = Path(filename).suffix.lower()
    _check_signature(data, suffix)
    document = Document(
        id=document_id,
        filename=Path(filename).name,
        revision=revision,
        sha256=hashlib.sha256(data).hexdigest(),
    )
    bundle = Bundle(document=document, media_type=suffix.lstrip("."))
    if suffix == ".pdf":
        from pdfminer.pdfdocument import PDFPasswordIncorrect

        try:
            pdf = pdfplumber.open(io.BytesIO(data))
        except PDFPasswordIncorrect:
            raise ValueError("PDF cifrado con contraseña: hay que pedir el original sin cifrar")
        except Exception as exc:  # pdfminer lanza sus propias excepciones con bytes corruptos
            raise ValueError(f"No es un PDF válido: {exc}") from None
        with pdf:
            if getattr(pdf.doc, "encryption", None):
                bundle.warnings.append(
                    "PDF cifrado sin contraseña de apertura: se lee, pero conviene el original"
                )
            if len(pdf.pages) > MAX_PAGES:
                raise ValueError("Límite de la base: 100 páginas por PDF")
            document.page_count = len(pdf.pages)
            for page in pdf.pages:
                clipped = dropped = 0
                page_lines: list[EvidenceLine] = []
                rows = page.extract_text_lines(return_chars=True)
                if not rows:
                    bundle.warnings.append(
                        f"Página {page.page_number}: sin texto; necesita OCR/visión"
                    )
                for row in rows:
                    # Separar columnas que comparten una misma línea base: en un plano, un hueco de
                    # más de 12 pt (unos cuatro espacios) es un cambio de celda o de bloque.
                    groups = [[]]
                    for char in sorted(row["chars"], key=lambda c: c["x0"]):
                        if groups[-1] and char["x0"] - groups[-1][-1]["x1"] > COLUMN_GAP_PT:
                            groups.append([])
                        groups[-1].append(char)
                    for group in groups:
                        parts = []
                        for i, char in enumerate(group):
                            if i and char["x0"] - group[i - 1]["x1"] > 1.0:
                                parts.append(" ")
                            parts.append(char["text"])
                        text = "".join(parts).strip()
                        if not text:
                            continue
                        raw = (
                            min(c["x0"] for c in group) / page.width,
                            min(c["top"] for c in group) / page.height,
                            max(c["x1"] for c in group) / page.width,
                            max(c["bottom"] for c in group) / page.height,
                        )
                        bbox = _clip_bbox(raw)
                        if bbox is None:
                            dropped += 1
                            continue
                        clipped += bbox != raw
                        page_lines.append(
                            EvidenceLine(
                                text=text,
                                source=Source(
                                    document_id=document_id,
                                    page=page.page_number,
                                    locator=f"Texto de página {page.page_number}",
                                    bbox=bbox,
                                    confidence=0.75,
                                ),
                            )
                        )
                bundle.lines.extend(_merge_rotated(page_lines))
                if dropped:
                    bundle.warnings.append(
                        f"Página {page.page_number}: {dropped} texto(s) fuera de la página, "
                        "descartado(s)"
                    )
                if clipped:
                    bundle.warnings.append(
                        f"Página {page.page_number}: {clipped} texto(s) recortado(s) al borde"
                    )
    elif suffix == ".xlsx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Libro demasiado grande al descomprimir")
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        try:
            total = 0
            for sheet in workbook:
                if sheet.max_row * sheet.max_column > 20000:
                    raise ValueError("Límite de la base: 20.000 celdas por hoja")
                for row in sheet.iter_rows():
                    cells = [c for c in row if c.value is not None]
                    total += len(cells)
                    if total > 50000:
                        raise ValueError("Límite de la base: 50.000 celdas por libro")
                    if not cells:
                        continue
                    if any(c.data_type == "f" for c in cells):
                        bundle.warnings.append(f"{sheet.title}!{cells[0].row}: fórmula sin evaluar")
                        continue
                    values = [str(c.value) for c in cells]
                    if len(values) > 2 and values[-1] == "1":
                        values = values[:-1]
                    text = f"{values[0]} = {' '.join(values[1:])}" if len(values) > 1 else values[0]
                    location = f"{cells[0].coordinate}:{cells[-1].coordinate}"
                    bundle.lines.append(
                        EvidenceLine(
                            text=text,
                            source=Source(
                                document_id=document_id,
                                sheet=sheet.title,
                                cell=location,
                                locator=f"{sheet.title}!{location}",
                                confidence=0.75,
                            ),
                        )
                    )
        finally:
            workbook.close()
    elif suffix in IMAGE_SUFFIXES:
        with Image.open(io.BytesIO(data)) as image:
            image.verify()
        with Image.open(io.BytesIO(data)) as image:
            document.page_count = getattr(image, "n_frames", 1)
        bundle.warnings.append(
            "Imagen recibida: `trazia.pipeline.process` la convierte en PDF escaneado para el OCR"
        )
    elif suffix in (".xls", ".xlsm"):
        raise ValueError(
            f"{suffix} no se admite: hay que guardarlo como .xlsx (las macros no se ejecutan nunca)"
        )
    else:
        raise ValueError("Formato admitido: PDF, XLSX, PNG, JPEG o TIFF")
    return bundle


COLUMN_GAP_PT = 12


def _merge_rotated(lines: list[EvidenceLine]) -> list[EvidenceLine]:
    """Une los trozos de un texto girado 90° que pdfplumber entrega apilados.

    Una cota girada como «559 1/16"» llega en dos líneas estrechas, una encima de otra, separadas
    por el espacio. Se unen de arriba abajo con un espacio: el resultado conserva el orden de
    caracteres de pdfplumber (invertido), igual que una cota girada de un solo trozo.
    """
    if not lines:
        return lines

    def vertical(line: EvidenceLine) -> bool:
        x0, y0, x1, y1 = line.source.bbox
        return (y1 - y0) > 1.5 * (x1 - x0)

    pending = sorted(lines, key=lambda line: (line.source.bbox[0], line.source.bbox[1]))
    out: list[EvidenceLine] = []
    for line in pending:
        last = out[-1] if out else None
        if last is not None:
            a0, b0, a1, b1 = last.source.bbox
            x0, y0, x1, y1 = line.source.bbox
            same_column = abs(a0 - x0) < 0.003 and abs(a1 - x1) < 0.003 and (a1 - a0) < 0.015
            if same_column and 0 <= y0 - b1 < 0.006 and (vertical(last) or vertical(line)):
                bbox = (min(a0, x0), b0, max(a1, x1), max(b1, y1))
                out[-1] = EvidenceLine(
                    text=f"{last.text} {line.text}",
                    source=last.source.model_copy(update={"bbox": bbox}),
                )
                continue
        out.append(line)
    return out


def _clip_bbox(raw: tuple[float, float, float, float]) -> tuple | None:
    """Recorta una caja normalizada a la página; None si queda fuera del todo."""
    x0, y0, x1, y1 = (min(1.0, max(0.0, v)) for v in raw)
    if x1 <= x0 or y1 <= y0:
        return None
    return (x0, y0, x1, y1)


def render_evidence(
    data: bytes, filename: str, page: int = 1, bbox: tuple | None = None
) -> Image.Image:
    if Path(filename).suffix.lower() == ".pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if not 1 <= page <= len(pdf.pages):
                raise ValueError("Página fuera del documento")
            selected = pdf.pages[page - 1]
            image = selected.to_image(resolution=110)
            if bbox:
                x0, y0, x1, y1 = bbox
                image.draw_rect(
                    (
                        x0 * selected.width,
                        y0 * selected.height,
                        x1 * selected.width,
                        y1 * selected.height,
                    ),
                    stroke="orange",
                    fill=None,
                    stroke_width=3,
                )
            return image.annotated.copy()
    with Image.open(io.BytesIO(data)) as image:
        return image.convert("RGB")
