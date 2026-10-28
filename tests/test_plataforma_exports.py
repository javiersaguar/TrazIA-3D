import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from trazia.cli import main
from trazia.model import Alternative, Interval, fields, load_revision
from trazia.review import approve, export_fields

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


def read_table(path):
    if path.suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    workbook = load_workbook(path, data_only=False)
    try:
        values = list(workbook.active.values)
        return [dict(zip(values[0], row, strict=True)) for row in values[1:]]
    finally:
        workbook.close()


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_field_table_preserves_review_evidence_and_uncertainty(tmp_path, extension):
    revision = load_revision(EXAMPLES / "vessel_rev_a.json")
    height = revision.equipment.height
    height.interval = Interval(low=5997.0, high=6003.0, level=0.9)
    height.state = "conflict"
    height.explanation = "Revisar la cota: diámetro y altura no son lo mismo."
    height.alternatives = [Alternative(value=6500.0, unit="mm", source=height.source)]
    before = revision.model_dump_json()
    target = export_fields(revision, tmp_path / f"campos.{extension}")
    rows = read_table(target)
    assert {r["field"] for r in rows} == set(fields(revision.equipment))
    row = next(r for r in rows if r["field"] == "height")
    source = height.source
    document = next(d for d in revision.documents if d.id == source.document_id)
    assert float(row["value"]) == height.value and row["unit"] == "mm"
    assert row["state"] == "conflict" and row["revision_status"] == "DRAFT"
    assert row["dataset"] == "synthetic"
    assert row["document_sha256"] == document.sha256
    assert row["document_id"] == source.document_id
    assert int(row["page"]) == source.page
    assert row["locator"] == source.locator
    assert json.loads(row["bbox"]) == list(source.bbox)
    assert float(row["interval_level"]) == 0.9
    assert float(row["interval_low"]) == 5997.0
    assert json.loads(row["alternatives"])[0]["value"] == 6500.0
    assert row["explanation"] == height.explanation
    assert revision.model_dump_json() == before


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_document_text_is_literal_in_spreadsheets(tmp_path, extension):
    revision = load_revision(EXAMPLES / "vessel_rev_a.json")
    revision.equipment.attributes["material"].value = '=HYPERLINK("https://invalid")'
    target = export_fields(revision, tmp_path / f"datos.{extension}")
    rows = read_table(target)
    value = next(r["value"] for r in rows if r["field"] == "attributes.material")
    if extension == "csv":
        assert value.startswith("'=HYPERLINK")
    else:
        assert value.startswith("=HYPERLINK")
        workbook = load_workbook(target, data_only=False)
        try:
            assert all(c.data_type != "f" for row in workbook.active for c in row)
        finally:
            workbook.close()


def test_export_preserves_approval_and_rejects_stale_in_memory_approval(tmp_path):
    revision = load_revision(EXAMPLES / "vessel_rev_a.json")
    originals = {
        d.id: (EXAMPLES / "documents" / d.filename).read_bytes() for d in revision.documents
    }
    approved = approve(revision, "Javier", originals)
    target = export_fields(approved, tmp_path / "aprobado.csv")
    rows = read_table(target)
    assert all(r["revision_status"] == "APPROVED" and r["reviewer"] == "Javier" for r in rows)
    approved.equipment.height.value += 1
    with pytest.raises(ValueError, match="aprobación"):
        export_fields(approved, tmp_path / "obsoleto.csv")
    assert not (tmp_path / "obsoleto.csv").exists()


def test_cli_exports_table_and_rejects_unknown_format(tmp_path):
    source = EXAMPLES / "vessel_rev_a.json"
    target = tmp_path / "nested" / "campos.xlsx"
    assert main(["fields", str(source), "-o", str(target)]) == 0
    assert read_table(target)
    assert main(["fields", str(source), "-o", str(tmp_path / "campos.pdf")]) == 2


def test_excel_does_not_silently_truncate_evidence(tmp_path):
    revision = load_revision(EXAMPLES / "vessel_rev_a.json")
    revision.equipment.height.source.text_original = "x" * 32768
    with pytest.raises(ValueError, match="límite de Excel"):
        export_fields(revision, tmp_path / "demasiado.xlsx")
    assert not (tmp_path / "demasiado.xlsx").exists()


def test_cli_markdown_diff_contains_values_actions_and_trace(tmp_path):
    before, after = (EXAMPLES / f"vessel_rev_{key}.json" for key in ("a", "b"))
    target = tmp_path / "cambios.md"
    assert main(["diff", str(before), str(after), "-o", str(target)]) == 0
    text = target.read_text(encoding="utf-8")
    assert "Datos sintéticos" in text and "Cambios de diseño: 8" in text
    assert "6000.0 mm" in text and "6500.0 mm" in text
    assert "Piping" in text and "Civil" in text
    assert "DRAFT" in text and "APPROVED" not in text
    for revision in (load_revision(before), load_revision(after)):
        assert revision.fingerprint() in text
        assert all(d.sha256 in text for d in revision.documents)


def test_diff_format_can_be_explicit_even_with_an_md_suffix(tmp_path, capsys):
    before, after = (str(EXAMPLES / f"vessel_rev_{key}.json") for key in ("a", "b"))
    target = tmp_path / "salida.md"
    assert main(["diff", before, after, "--format", "json", "-o", str(target)]) == 0
    assert json.loads(target.read_text(encoding="utf-8"))["summary"]["changes"] == 8
    assert main(["diff", before, after, "--format", "markdown"]) == 0
    assert "# Informe de cambios" in capsys.readouterr().out


def test_markdown_keeps_document_text_inside_the_table():
    from trazia.revisions import compare, markdown_report

    before = load_revision(EXAMPLES / "vessel_rev_a.json")
    after = before.model_copy(deep=True)
    after.equipment.attributes["material"].value = "<script> | fila\nnueva"
    report = markdown_report(before, after, compare(before, after))
    assert "<script>" not in report and "&lt;script&gt;" in report
    assert "\\| fila nueva" in report
