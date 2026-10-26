"""Generar documentos sintéticos, ingerirlos y construir revisiones DRAFT reproducibles."""

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A3, landscape
from reportlab.pdfgen.canvas import Canvas

from trazia.extraction import LabelExtractor, assemble
from trazia.ingest import ingest_bytes

ROOT = Path(__file__).resolve().parents[1]


def make_ga(path, revision, values, nozzles):
    canvas = Canvas(str(path), pagesize=landscape(A3), invariant=1)
    canvas.setTitle(f"TrazIA synthetic GA V-101 rev {revision}")
    navy, teal = HexColor("#172c3d"), HexColor("#007e80")
    canvas.setFillColor(navy)
    canvas.setFont("Helvetica-Bold", 26)
    canvas.drawString(40, 790, f"TrazIA / V-101 / REV {revision}")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(40, 766, "SYNTHETIC FIXTURE - Not a vendor document - All lengths in mm")
    canvas.setStrokeColor(navy)
    canvas.setFillColor(HexColor("#edf4f6"))
    x, bottom, width, shell = 160, 180, 170, 390
    head = values["head_depth"] / values["height"] * shell
    support = values["skirt_height"] / values["height"] * shell
    canvas.rect(x, bottom, width, shell, fill=1)
    for z, sign in ((bottom, -1), (bottom + shell, 1)):
        p = canvas.beginPath()
        p.moveTo(x, z)
        p.curveTo(x, z + sign * head * 1.333, x + width, z + sign * head * 1.333, x + width, z)
        p.close()
        canvas.drawPath(p, stroke=1, fill=1)
    canvas.setFillColor(HexColor("#d2e0e7"))
    canvas.rect(x - 4, bottom - support, 9, support, fill=1)
    canvas.rect(x + width - 5, bottom - support, 9, support, fill=1)
    canvas.line(90, bottom - support, 420, bottom - support)
    canvas.setFillColor(navy)
    canvas.setFont("Helvetica", 10)
    canvas.drawString(92, bottom - support - 18, "Z = 0 / ground")
    canvas.drawString(x + 42, bottom + shell / 2, "V-101")
    canvas.drawString(100, 650, "SIMPLIFIED ELEVATION / not to scale")
    for i, (name, diameter, projection, elevation, azimuth) in enumerate(nozzles):
        z = bottom + (elevation - values["skirt_height"]) / values["height"] * shell
        cosine = math.cos(math.radians(azimuth))
        if abs(cosine) < 0.1:
            canvas.setStrokeColor(teal)
            canvas.setFillColor(teal)
            canvas.circle(x + width / 2, z, 5, fill=0)
            canvas.drawString(x + width / 2 + 12, z + 8, f"{name} ({azimuth} deg)")
            continue
        right = cosine > 0
        start = x + width if right else x
        end = start + (45 if right else -45)
        canvas.setStrokeColor(teal)
        canvas.line(start, z, end, z)
        canvas.setFillColor(teal)
        canvas.drawString(end - (0 if right else 20), z + 8, name)
    canvas.setStrokeColor(navy)
    canvas.setFillColor(navy)
    canvas.setFont("Helvetica-Bold", 15)
    canvas.drawString(510, 714, "EXPLICIT PARAMETERS / EXTRACTION BASELINE")
    canvas.setFont("Helvetica", 12)
    canvas.drawString(510, 682, "TAG = V-101")
    for i, (name, value) in enumerate(values.items()):
        canvas.drawString(510, 650 - 27 * i, f"{name.upper()} = {value} mm")
    canvas.setFont("Helvetica-Bold", 15)
    canvas.drawString(510, 474, "NOZZLE SCHEDULE / RADIAL CONNECTIONS")
    canvas.setFont("Helvetica", 9)
    for i, (name, diameter, projection, elevation, azimuth) in enumerate(nozzles):
        line = (
            f"NOZZLE {name}: DIAMETER={diameter} mm; PROJECTION={projection} mm; "
            f"ELEVATION={elevation} mm; AZIMUTH={azimuth} deg"
        )
        canvas.drawString(510, 443 - 32 * i, line)
    canvas.setFont("Helvetica", 11)
    notes = [
        "DATUMS / CONVENTIONS",
        "HEIGHT: straight shell length between tangent lines.",
        "SKIRT_HEIGHT: ground to lower tangent line.",
        "HEAD_DEPTH: equal semi-ellipsoidal heads, upper and lower.",
        "ELEVATION: nozzle axis above ground. PROJECTION: from shell surface.",
        "AZIMUTH: counter-clockwise from +X toward +Y.",
        "Document for software tests; no pressure design or fabrication validity.",
    ]
    for i, line in enumerate(notes):
        canvas.drawString(510, 254 - 22 * i, line)
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawString(40, 45, "TRAZIA-3D | GA2MODEL | SYNTHETIC ENGINEERING DATA")
    canvas.save()


def make_sheet(path, weight):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "HDM"
    rows = [
        ("PARAMETER", "VALUE", "UNIT"),
        ("DESIGN_PRESSURE", 16, "bar"),
        ("DESIGN_TEMPERATURE", 180, "C"),
        ("MATERIAL", "SA-516-70", "1"),
        ("OPERATING_WEIGHT", weight, "kg"),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="172C3D")
    for column, width in (("A", 28), ("B", 22), ("C", 12)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)


def main():
    directory = ROOT / "examples" / "documents"
    directory.mkdir(parents=True, exist_ok=True)
    for revision in ("A", "B"):
        second = revision == "B"
        values = {
            "diameter": 2000,
            "height": 6500 if second else 6000,
            "head_depth": 500,
            "skirt_height": 1350 if second else 1200,
            "skirt_thickness": 20,
        }
        nozzles = [
            ("N1", 200, 400, 2500, 0),
            ("N2", 100, 300, 5500, 90),
            ("N3", 150, 350, 3700 if second else 3500, 180),
        ]
        if second:
            nozzles.append(("N4", 100, 300, 6200, 270))
        ga = directory / f"GAD_V-101_rev{revision}.pdf"
        hdm = directory / f"HDM_V-101_rev{revision}.xlsx"
        make_ga(ga, revision, values, nozzles)
        make_sheet(hdm, 12800 if second else 12000)
        bundles = [
            ingest_bytes(p.read_bytes(), p.name, name, revision)
            for p, name in ((ga, "GA-V101"), (hdm, "HDM-V101"))
        ]
        candidates = LabelExtractor().extract(bundles)
        model = assemble(
            bundles,
            candidates,
            project_id="DEMO-PLANT",
            equipment_id="vessel-001",
            revision=revision,
            dataset="synthetic",
        )
        (ROOT / "examples" / f"vessel_rev_{revision.lower()}.json").write_text(
            model.model_dump_json(indent=2) + "\n", encoding="utf-8"
        )
        print(
            f"Revision {revision}: {len(candidates)} candidatos, "
            f"{len(model.equipment.nozzles)} boquillas"
        )


if __name__ == "__main__":
    main()
