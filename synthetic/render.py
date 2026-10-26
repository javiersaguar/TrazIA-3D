"""Documentos sintéticos (GA en PDF y hoja de datos en Excel) en tres estilos.

- rotulado: parámetros escritos como «ETIQUETA = valor unidad» (lo que lee el baseline actual).
- plano: GA dibujado como un plano real: alzado acotado, planta con orientaciones, tabla de
  boquillas, datos de diseño y cajetín. Las cotas no llevan unidad (nota «mm salvo indicación»).
- escaneado: el plano rasterizado con ruido y una leve rotación, sin capa de texto.
"""

import io
import math
import random
import zlib
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageFilter
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A3, landscape
from reportlab.pdfgen.canvas import Canvas

from .cases import VesselSpec

INK = HexColor("#1b2a38")
FAINT = HexColor("#8a99a8")
ACCENT = HexColor("#0b6e8f")
REVISION = HexColor("#c2410c")

STYLES = ("rotulado", "plano", "escaneado")


def write_hdm(path: Path, spec: VesselSpec) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "HDM"
    rows = [
        ("PARAMETER", "VALUE", "UNIT"),
        ("DESIGN_PRESSURE", spec.design_pressure, "bar"),
        ("DESIGN_TEMPERATURE", spec.design_temperature, "C"),
        ("MATERIAL", spec.material, "1"),
        ("OPERATING_WEIGHT", spec.operating_weight, "kg"),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B2A38")
    for column, width in (("A", 28), ("B", 22), ("C", 12)):
        sheet.column_dimensions[column].width = width
    workbook.save(path)


def write_ga(
    path: Path, spec: VesselSpec, revision: str, style: str, changed: set[str] | None = None
) -> None:
    if style == "rotulado":
        _ga_labelled(path, spec, revision)
    elif style == "plano":
        _ga_drawing(path, spec, revision, changed or set())
    elif style == "escaneado":
        buffer = io.BytesIO()
        _ga_drawing(buffer, spec, revision, changed or set())
        _scan(buffer.getvalue(), path, seed=zlib.crc32(f"{spec.tag}/{revision}".encode()))
    else:
        raise ValueError(f"Estilo desconocido: {style}")


def _canvas(target, spec: VesselSpec, revision: str) -> Canvas:
    canvas = Canvas(
        target if isinstance(target, io.BytesIO) else str(target),
        pagesize=landscape(A3),
        invariant=1,
    )
    canvas.setTitle(f"Synthetic GA {spec.tag} rev {revision}")
    return canvas


def _ga_labelled(path: Path, spec: VesselSpec, revision: str) -> None:
    canvas = _canvas(path, spec, revision)
    canvas.setFillColor(INK)
    canvas.setFont("Helvetica-Bold", 24)
    canvas.drawString(40, 790, f"SYNTHETIC GA / {spec.tag} / REV {revision}")
    canvas.setFont("Helvetica", 11)
    canvas.drawString(40, 766, "SYNTHETIC FIXTURE - Not a vendor document - All lengths in mm")
    _elevation(canvas, spec, x=150, base=150, width=170, shell=_fit_shell(spec, 520), dims=False)
    canvas.setFillColor(INK)
    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(520, 720, "EXPLICIT PARAMETERS")
    canvas.setFont("Helvetica", 12)
    canvas.drawString(520, 690, f"TAG = {spec.tag}")
    values = {
        "DIAMETER": spec.diameter,
        "HEIGHT": spec.height,
        "HEAD_DEPTH": spec.head_depth,
        "SKIRT_HEIGHT": spec.skirt_height,
        "SKIRT_THICKNESS": spec.skirt_thickness,
    }
    for i, (name, value) in enumerate(values.items()):
        canvas.drawString(520, 660 - 26 * i, f"{name} = {value} mm")
    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(520, 500, "NOZZLE SCHEDULE")
    canvas.setFont("Helvetica", 9)
    for i, n in enumerate(spec.nozzles):
        canvas.drawString(
            520,
            474 - 24 * i,
            f"NOZZLE {n.id}: DIAMETER={n.diameter} mm; PROJECTION={n.projection} mm; "
            f"ELEVATION={n.elevation} mm; AZIMUTH={n.azimuth} deg",
        )
    canvas.setFont("Helvetica-Bold", 9)
    canvas.drawString(40, 40, "TRAZIA-3D | GA2MODEL | SYNTHETIC ENGINEERING DATA")
    canvas.save()


def _elevation(
    canvas: Canvas,
    spec: VesselSpec,
    *,
    x: float,
    base: float,
    width: float,
    shell: float,
    dims: bool,
) -> dict:
    """Alzado esquemático a escala vertical. Devuelve la escala para posicionar boquillas."""
    scale = shell / spec.height
    head = spec.head_depth * scale
    skirt = spec.skirt_height * scale
    bottom = base + skirt
    canvas.setStrokeColor(INK)
    canvas.setLineWidth(1)
    canvas.setFillColor(HexColor("#f1f5f8"))
    canvas.rect(x, bottom, width, shell, fill=1)
    for z, sign in ((bottom, -1), (bottom + shell, 1)):
        p = canvas.beginPath()
        p.moveTo(x, z)
        p.curveTo(x, z + sign * head * 1.333, x + width, z + sign * head * 1.333, x + width, z)
        p.close()
        canvas.drawPath(p, stroke=1, fill=1)
    canvas.setFillColor(HexColor("#dfe6ec"))
    canvas.rect(x + 6, base, 8, skirt, fill=1)
    canvas.rect(x + width - 14, base, 8, skirt, fill=1)
    canvas.line(x - 60, base, x + width + 60, base)
    canvas.setDash(6, 3)
    canvas.setStrokeColor(FAINT)
    canvas.line(x + width / 2, base - 10, x + width / 2, bottom + shell + head + 20)
    canvas.line(x - 30, bottom, x + width + 30, bottom)
    canvas.line(x - 30, bottom + shell, x + width + 30, bottom + shell)
    canvas.setDash()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(FAINT)
    canvas.drawString(x + width + 34, bottom - 3, "T.L.")
    canvas.drawString(x + width + 34, bottom + shell - 3, "T.L.")
    for n in spec.nozzles:
        z = base + n.elevation * scale
        c = math.cos(math.radians(n.azimuth))
        canvas.setStrokeColor(ACCENT)
        canvas.setFillColor(ACCENT)
        if abs(c) < 0.3:
            canvas.circle(x + width / 2 + c * width * 0.4, z, 4, fill=0)
            canvas.setFont("Helvetica", 8)
            canvas.drawString(x + width / 2 + c * width * 0.4 + 7, z + 4, n.id)
            continue
        start = x + width if c > 0 else x
        end = start + (36 if c > 0 else -36)
        canvas.line(start, z, end, z)
        canvas.line(end, z - 5, end, z + 5)
        canvas.setFont("Helvetica", 8)
        canvas.drawString(end + (4 if c > 0 else -18), z + 4, n.id)
    return {"scale": scale, "bottom": bottom, "base": base, "head": head}


def _fit_shell(spec: VesselSpec, budget: float) -> float:
    """Altura de dibujo de la virola para que faldón + virola + fondo quepan en «budget» puntos."""
    total = spec.skirt_height + spec.height + spec.head_depth * 1.333
    return spec.height * budget / total


def _dimension_v(canvas: Canvas, x: float, z0: float, z1: float, text: str, side: int = -1) -> None:
    """Cota vertical: líneas de extensión, línea de cota con flechas y texto girado."""
    canvas.setStrokeColor(INK)
    canvas.setFillColor(INK)
    canvas.setLineWidth(0.6)
    canvas.line(x, z0, x, z1)
    for z, d in ((z0, 1), (z1, -1)):
        canvas.line(x - 3, z + 6 * d, x, z)
        canvas.line(x + 3, z + 6 * d, x, z)
    canvas.saveState()
    canvas.translate(x + 4 * side - (8 if side < 0 else -2), (z0 + z1) / 2)
    canvas.rotate(90)
    canvas.setFont("Helvetica", 9)
    canvas.drawCentredString(0, 0, text)
    canvas.restoreState()


def _delta(canvas: Canvas, x: float, y: float, revision: str) -> None:
    """Triángulo de revisión, como en los planos reales."""
    canvas.setStrokeColor(REVISION)
    canvas.setFillColor(REVISION)
    p = canvas.beginPath()
    p.moveTo(x, y)
    p.lineTo(x + 12, y)
    p.lineTo(x + 6, y + 10)
    p.close()
    canvas.drawPath(p, stroke=1, fill=0)
    canvas.setFont("Helvetica-Bold", 6)
    canvas.drawCentredString(x + 6, y + 2, revision)


def _ga_drawing(target, spec: VesselSpec, revision: str, changed: set[str]) -> None:
    canvas = _canvas(target, spec, revision)
    w, h = landscape(A3)
    canvas.setStrokeColor(INK)
    canvas.setLineWidth(1.2)
    canvas.rect(20, 20, w - 40, h - 40)

    # Alzado con cotas
    geo = _elevation(
        canvas, spec, x=200, base=110, width=150, shell=_fit_shell(spec, 600), dims=True
    )
    base, bottom, scale, head = geo["base"], geo["bottom"], geo["scale"], geo["head"]
    top = bottom + spec.height * scale
    _dimension_v(canvas, 120, base, bottom, f"{spec.skirt_height}")
    _dimension_v(canvas, 120, bottom, top, f"{spec.height}")
    _dimension_v(
        canvas, 80, base, top + head, f"{spec.skirt_height + spec.height + spec.head_depth}"
    )
    _dimension_v(canvas, 160, top, top + head, f"{spec.head_depth}")
    canvas.setLineWidth(0.6)
    canvas.line(200, top + head + 40, 350, top + head + 40)
    canvas.setFont("Helvetica", 9)
    canvas.drawCentredString(275, top + head + 46, f"ID {spec.diameter}")
    canvas.drawCentredString(275, base - 18, f"SKIRT t={spec.skirt_thickness}")
    if changed & {"height", "skirt_height"}:
        _delta(canvas, 100, (bottom + top) / 2, revision)
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawString(150, 70, "ELEVATION")

    # Planta con orientaciones
    cx, cy, r = 520, 560, 90
    canvas.setStrokeColor(INK)
    canvas.circle(cx, cy, r)
    canvas.setStrokeColor(FAINT)
    canvas.setDash(4, 3)
    canvas.line(cx - r - 30, cy, cx + r + 30, cy)
    canvas.line(cx, cy - r - 30, cx, cy + r + 30)
    canvas.setDash()
    canvas.setFillColor(INK)
    canvas.setFont("Helvetica", 8)
    for angle, label in ((0, "0°"), (90, "90°"), (180, "180°"), (270, "270°")):
        a = math.radians(angle)
        canvas.drawCentredString(
            cx + (r + 42) * math.cos(a), cy + (r + 42) * math.sin(a) - 3, label
        )
    for n in spec.nozzles:
        a = math.radians(n.azimuth)
        x0, y0 = cx + r * math.cos(a), cy + r * math.sin(a)
        x1, y1 = cx + (r + 22) * math.cos(a), cy + (r + 22) * math.sin(a)
        canvas.setStrokeColor(ACCENT)
        canvas.line(x0, y0, x1, y1)
        canvas.setFillColor(ACCENT)
        canvas.drawCentredString(cx + (r + 30) * math.cos(a), cy + (r + 30) * math.sin(a) + 6, n.id)
    canvas.setFillColor(INK)
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawString(470, 420, "PLAN · ORIENTATION")

    # Tabla de boquillas
    tx, ty = 700, 760
    cols = (
        ("MARK", 50),
        ("SIZE DN", 70),
        ("PROJ.", 70),
        ("ELEV.", 80),
        ("ORIENT.", 70),
        ("SERVICE", 120),
    )
    canvas.setFont("Helvetica-Bold", 11)
    canvas.drawString(tx, ty + 14, "NOZZLE SCHEDULE")
    row_h = 18
    total_w = sum(c[1] for c in cols)
    canvas.setStrokeColor(INK)
    canvas.setLineWidth(0.6)
    canvas.rect(tx, ty - row_h * (len(spec.nozzles) + 1), total_w, row_h * (len(spec.nozzles) + 1))
    x = tx
    canvas.setFont("Helvetica-Bold", 8)
    for name, width in cols:
        canvas.drawString(x + 4, ty - 12, name)
        x += width
        canvas.line(x, ty, x, ty - row_h * (len(spec.nozzles) + 1))
    services = ("INLET", "OUTLET", "VENT", "DRAIN", "INSTR.", "MANWAY", "SPARE", "RELIEF")
    canvas.setFont("Helvetica", 8)
    for i, n in enumerate(spec.nozzles):
        y = ty - row_h * (i + 1)
        canvas.line(tx, y, tx + total_w, y)
        values = (
            n.id,
            f"{n.diameter}",
            f"{n.projection}",
            f"{n.elevation}",
            f"{n.azimuth}°",
            services[i % len(services)],
        )
        x = tx
        for (_, width), value in zip(cols, values, strict=True):
            canvas.drawString(x + 4, y - 12, value)
            x += width
        if any(p.startswith(f"nozzles.{n.id}.") for p in changed):
            _delta(canvas, tx + total_w + 6, y - 14, revision)

    # Datos de diseño (sin etiquetas «=»: solo tabla)
    dy = ty - row_h * (len(spec.nozzles) + 1) - 50
    canvas.setFont("Helvetica-Bold", 11)
    canvas.drawString(tx, dy + 14, "DESIGN DATA")
    data = (
        ("DESIGN PRESSURE", f"{spec.design_pressure} barg", "attributes.design_pressure"),
        ("DESIGN TEMPERATURE", f"{spec.design_temperature} °C", "attributes.design_temperature"),
        ("SHELL MATERIAL", spec.material, "attributes.material"),
        ("OPERATING WEIGHT", f"{spec.operating_weight} kg", "attributes.operating_weight"),
        ("HEADS", "2:1 SEMI-ELLIPSOIDAL", "head_depth"),
    )
    for i, (k, v, path) in enumerate(data):
        y = dy - 16 * i
        canvas.setStrokeColor(INK)
        canvas.setFillColor(INK)
        canvas.setFont("Helvetica", 8)
        canvas.line(tx, y - 4, tx + 300, y - 4)
        canvas.drawString(tx + 4, y, k)
        canvas.drawString(tx + 170, y, v)
        if path in changed:
            _delta(canvas, tx + 306, y - 2, revision)

    # Notas y cajetín (anonimizado, como el dataset del reto)
    canvas.setFont("Helvetica", 8)
    notes = (
        "NOTES",
        "1. ALL DIMENSIONS IN MM UNLESS OTHERWISE NOTED.",
        "2. ELEVATIONS FROM GROUND (Z = 0). T.L. = TANGENT LINE.",
        "3. ORIENTATION COUNTER-CLOCKWISE FROM 0°.",
        "4. SYNTHETIC DOCUMENT FOR SOFTWARE TESTS.",
    )
    for i, line in enumerate(notes):
        canvas.drawString(tx, 230 - 13 * i, line)
    canvas.rect(w - 360, 30, 330, 90)
    canvas.line(w - 360, 75, w - 30, 75)
    canvas.line(w - 200, 30, w - 200, 120)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(FAINT)
    canvas.drawString(w - 355, 110, "EQUIPMENT")
    canvas.drawString(w - 195, 110, "DOCUMENT")
    canvas.drawString(w - 355, 65, "TITLE")
    canvas.drawString(w - 195, 65, "REV")
    canvas.setFillColor(INK)
    canvas.setFont("Helvetica-Bold", 12)
    canvas.drawString(w - 355, 88, spec.tag)
    canvas.drawString(w - 195, 88, f"GAD-{spec.tag}")
    canvas.setFont("Helvetica", 9)
    canvas.drawString(w - 355, 42, "GENERAL ARRANGEMENT")
    canvas.setFont("Helvetica-Bold", 14)
    canvas.drawString(w - 195, 42, revision)
    canvas.save()


def _scan(pdf_bytes: bytes, path: Path, seed: int) -> None:
    rng = random.Random(seed)
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        image = pdf.pages[0].to_image(resolution=110).original.convert("L")
    image = image.rotate(
        rng.uniform(-0.8, 0.8), resample=Image.Resampling.BICUBIC, expand=False, fillcolor=255
    )
    image = image.filter(ImageFilter.GaussianBlur(0.6))
    noise = Image.effect_noise(image.size, 18).convert("L")
    image = Image.blend(image, noise, 0.08)
    image.save(path, "PDF", resolution=110)
